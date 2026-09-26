#!/usr/bin/env python3
"""
JAMB UTME Questions to Word (.docx) Exporter
Generates professional, teacher/mock-exam formatted Word documents per subject:
- Demarcated by exam year (recent first: 2025 -> 1978)
- Questions numbered 1 to N per year
- High-resolution embedded diagrams/illustrations
- Indented multiple choice options (A, B, C, D, E)
- No answers shown in question body
- Compact, beautifully formatted Answer Key table at the end of each year
- Page break between years
"""

import sys
import os
import json
import re
import argparse
from collections import defaultdict
from PIL import Image

import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

OPTIONS_LETTERS = ['A', 'B', 'C', 'D', 'E', 'F', 'G']

def clean_text(text):
    """
    Cleans LaTeX, HTML tags, and entities from question/option text
    converting them into clean, publication-ready Unicode characters.
    """
    if not text:
        return ""
    
    # HTML entities
    text = text.replace('&nbsp;', ' ').replace('&amp;', '&').replace('&lt;', '<')
    text = text.replace('&gt;', '>').replace('&quot;', '"').replace('&#39;', "'")
    
    # Basic HTML line breaks and paragraph tags
    text = re.sub(r'<br\s*/?>', '\n', text, flags=re.IGNORECASE)
    text = re.sub(r'</p\s*>', '\n', text, flags=re.IGNORECASE)
    text = re.sub(r'<[^>]+>', '', text)
    
    # Nigerian Currency: Naira symbol \N
    text = re.sub(r'\\N(?=[0-9])', '₦', text)
    text = text.replace(r'\N', '₦')
    
    # Common fractions
    frac_map = {
        r'\\frac\{1\}\{2\}': '½', r'\\frac\{1\}\{4\}': '¼', r'\\frac\{3\}\{4\}': '¾',
        r'\\frac\{1\}\{3\}': '⅓', r'\\frac\{2\}\{3\}': '⅔', r'\\frac\{1\}\{8\}': '⅛',
        r'\\frac\{3\}\{8\}': '⅜', r'\\frac\{5\}\{8\}': '⅝', r'\\frac\{7\}\{8\}': '⅞'
    }
    for pat, rep in frac_map.items():
        text = re.sub(pat, rep, text)
        
    # Square root \sqrt{...} or \sqrt ...
    text = re.sub(r'\\sqrt\{([^{}]+)\}', r'√(\1)', text)
    text = re.sub(r'\\sqrt([0-9a-zA-Z])', r'√\1', text)
    
    # Generic fraction \frac{a}{b} -> (a/b)
    for _ in range(3):
        text = re.sub(r'\\frac\{([^{}]+)\}\{([^{}]+)\}', r'(\1/\2)', text)

    # Subscripts mapping
    sub_map = {
        '0': '₀', '1': '₁', '2': '₂', '3': '₃', '4': '₄',
        '5': '₅', '6': '₆', '7': '₇', '8': '₈', '9': '₉',
        '+': '₊', '-': '₋', '=': '₌', '(': '₍', ')': '₎',
        'a': 'ₐ', 'e': 'ₑ', 'o': 'ₒ', 'x': 'ₓ', 'h': 'ₕ',
        'k': 'ₖ', 'l': 'ₗ', 'm': 'ₘ', 'n': 'ₙ', 'p': 'ₚ',
        's': 'ₛ', 't': 'ₜ', 'w': 'ᵥ'
    }
    # Superscripts mapping
    sup_map = {
        '0': '⁰', '1': '¹', '2': '²', '3': '³', '4': '⁴',
        '5': '⁵', '6': '⁶', '7': '⁷', '8': '⁸', '9': '⁹',
        '+': '⁺', '-': '⁻', '=': '⁼', '(': '⁽', ')': '⁾',
        'n': 'ⁿ', 'i': 'ⁱ', 'o': '°', 'x': 'ˣ', 'y': 'ʸ', 'z': 'ᶻ'
    }

    def rep_sub(m):
        return ''.join(sub_map.get(c, c) for c in m.group(1))

    def rep_sup(m):
        return ''.join(sup_map.get(c, c) for c in m.group(1))

    # LaTeX subscripts: _{...} or _0-9
    text = re.sub(r'_\{([a-zA-Z0-9+\-=()]+)\}', rep_sub, text)
    text = re.sub(r'_([0-9])', rep_sub, text)

    # LaTeX superscripts: ^{...} or ^0-9
    text = re.sub(r'\^\{([a-zA-Z0-9+\-=()]+)\}', rep_sup, text)
    text = re.sub(r'\^([0-9])', rep_sup, text)

    # Matrices & Arrays
    text = re.sub(r'\\begin\{(?:pmatrix|matrix|bmatrix)\}', '[ ', text)
    text = re.sub(r'\\end\{(?:pmatrix|matrix|bmatrix)\}', ' ]', text)
    text = re.sub(r'\\begin\{array\}\{[^}]*\}', '[ ', text)
    text = re.sub(r'\\end\{array\}', ' ]', text)
    text = re.sub(r'\\\\', ' ; ', text)
    text = re.sub(r'&', '  ', text)

    # Integrals, limits, sums
    text = text.replace(r'\int', '∫').replace(r'\sum', '∑').replace(r'\lim', 'lim')
    text = text.replace(r'\limits', '')

    # Delimiters
    text = re.sub(r'\\left\(', '(', text)
    text = re.sub(r'\\right\)', ')', text)
    text = re.sub(r'\\left\[', '[', text)
    text = re.sub(r'\\right\]', ']', text)
    text = re.sub(r'\\left\|', '|', text)
    text = re.sub(r'\\right\|', '|', text)
    text = re.sub(r'\\left\.', '', text)
    text = re.sub(r'\\right\.', '', text)

    # Arrows & Math Symbols
    sym_map = {
        r'\\longrightarrow': '→', r'\\rightarrow': '→', r'\\to': '→',
        r'\\rightleftharpoons': '⇌', r'\\leftarrow': '←', r'\\leftrightarrow': '↔',
        r'\\Rightarrow': '⇒', r'\\pm': '±', r'\\times': '×', r'\\div': '÷',
        r'\\leq': '≤', r'\\geq': '≥', r'\\neq': '≠', r'\\approx': '≈',
        r'\\circ': '°', r'\\degree': '°', r'\\alpha': 'α', r'\\beta': 'β',
        r'\\gamma': 'γ', r'\\theta': 'θ', r'\\lambda': 'λ', r'\\mu': 'μ',
        r'\\pi': 'π', r'\\sigma': 'σ', r'\\omega': 'ω', r'\\phi': 'φ',
        r'\\Delta': 'Δ', r'\\bigtriangleup': 'Δ', r'\\Omega': 'Ω', r'\\delta': 'δ',
        r'\\cap': '∩', r'\\cup': '∪', r'\\in': '∈', r'\\subseteq': '⊆',
        r'\\emptyset': '∅', r'\\oplus': '⊕', r'\\otimes': '⊗', r'\\angle': '∠',
        r'\\equiv': '≡', r'\\iff': '⟺', r'\\ast': '*', r'\\sin': 'sin',
        r'\\cos': 'cos', r'\\tan': 'tan', r'\\sec': 'sec', r'\\csc': 'csc',
        r'\\cot': 'cot', r'\\log': 'log', r'\\hline': '',
        r'----->': '→', r'--->': '→', r'-->': '→'
    }
    for pat, rep in sym_map.items():
        text = re.sub(pat, rep, text)

    # Strip text wrappers like \text{...}, \mathrm{...}
    text = re.sub(r'\\hspace\{[^}]*\}', ' ', text)
    text = re.sub(r'\\(?:text|mathrm|mathbf)\{([^{}]*)\}', r'\1', text)

    # Remove LaTeX enclosures \( \) \[ \] $
    text = text.replace(r'\(', '').replace(r'\)', '')
    text = text.replace(r'\[', '').replace(r'\]', '')
    text = re.sub(r'(?<!\\)\$', '', text)

    # Clean redundant whitespace
    text = re.sub(r'[ \t]+', ' ', text)
    return text.strip()


def set_cell_shading(cell, fill_hex):
    """Set background color of a table cell."""
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def set_table_borders(table, color="CBD5E0"):
    """Set thin, clean borders for a Word table."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:left w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
            <w:insideV w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)


def format_run(run, font_name="Calibri", size_pt=10.5, bold=False, italic=False, color_rgb=None):
    """Helper to consistently format text runs."""
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color_rgb:
        run.font.color.rgb = color_rgb


def add_answer_key_table(doc, year_questions, subject_title, year):
    """
    Creates a compact 10-column table (5 pairs of Q# / Answer) for the year's questions.
    """
    ak_head = doc.add_paragraph()
    ak_head.paragraph_format.space_before = Pt(16)
    ak_head.paragraph_format.space_after = Pt(6)
    ak_head.paragraph_format.keep_with_next = True
    
    ak_run = ak_head.add_run(f"ANSWER KEY — {subject_title.upper()} {year}")
    format_run(ak_run, font_name="Calibri", size_pt=12, bold=True, color_rgb=RGBColor(0x1A, 0x36, 0x5D))
    
    N = len(year_questions)
    NUM_PAIRS = 5  # 5 pairs = 10 columns
    num_data_rows = (N + NUM_PAIRS - 1) // NUM_PAIRS
    
    table = doc.add_table(rows=num_data_rows + 1, cols=NUM_PAIRS * 2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, color="CBD5E0")
    
    # Widths: Q# col = 0.45 in, Ans col = 0.85 in
    col_widths = [Inches(0.45), Inches(0.85)] * NUM_PAIRS
    
    # Format Header Row
    hdr_row = table.rows[0]
    trPr = hdr_row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    
    for c_idx in range(NUM_PAIRS * 2):
        cell = hdr_row.cells[c_idx]
        cell.width = col_widths[c_idx]
        set_cell_shading(cell, "1A365D")  # Deep navy
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        
        is_q_col = (c_idx % 2 == 0)
        label = "Q#" if is_q_col else "Ans"
        hrun = p.add_run(label)
        format_run(hrun, font_name="Calibri", size_pt=9.5, bold=True, color_rgb=RGBColor(0xFF, 0xFF, 0xFF))
        
    # Populate Data Cells
    for q_idx, q in enumerate(year_questions):
        q_num = q_idx + 1
        r_idx = (q_idx // NUM_PAIRS) + 1  # 1-indexed for header
        pair_idx = q_idx % NUM_PAIRS
        c_q = pair_idx * 2
        c_ans = pair_idx * 2 + 1
        
        correct_idx = q.get('correct')
        if isinstance(correct_idx, int) and 0 <= correct_idx < len(OPTIONS_LETTERS):
            ans_letter = OPTIONS_LETTERS[correct_idx]
        elif isinstance(correct_idx, str) and correct_idx.strip():
            ans_letter = correct_idx.strip().upper()
        else:
            ans_letter = "—"
            
        # Q cell
        cell_q = table.cell(r_idx, c_q)
        cell_q.width = col_widths[c_q]
        cell_q.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p_q = cell_q.paragraphs[0]
        p_q.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        p_q.paragraph_format.space_before = Pt(1.5)
        p_q.paragraph_format.space_after = Pt(1.5)
        run_q = p_q.add_run(str(q_num))
        format_run(run_q, font_name="Calibri", size_pt=9.5, color_rgb=RGBColor(0x4A, 0x55, 0x68))
        
        # Ans cell
        cell_ans = table.cell(r_idx, c_ans)
        cell_ans.width = col_widths[c_ans]
        cell_ans.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p_ans = cell_ans.paragraphs[0]
        p_ans.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
        p_ans.paragraph_format.space_before = Pt(1.5)
        p_ans.paragraph_format.space_after = Pt(1.5)
        run_ans = p_ans.add_run(ans_letter)
        format_run(run_ans, font_name="Calibri", size_pt=10, bold=True, color_rgb=RGBColor(0x1A, 0x36, 0x5D))
        
        # Row zebra shading
        if r_idx % 2 == 0:
            set_cell_shading(cell_q, "F7FAFC")
            set_cell_shading(cell_ans, "F7FAFC")

    # Format any remaining empty cells in the last row
    for q_idx in range(N, num_data_rows * NUM_PAIRS):
        r_idx = (q_idx // NUM_PAIRS) + 1
        pair_idx = q_idx % NUM_PAIRS
        c_q = pair_idx * 2
        c_ans = pair_idx * 2 + 1
        cell_q = table.cell(r_idx, c_q)
        cell_ans = table.cell(r_idx, c_ans)
        cell_q.width = col_widths[c_q]
        cell_ans.width = col_widths[c_ans]
        if r_idx % 2 == 0:
            set_cell_shading(cell_q, "F7FAFC")
            set_cell_shading(cell_ans, "F7FAFC")


def export_subject_docx(subject, output_dir="exports"):
    """
    Exports a single subject's past questions to Word (.docx).
    """
    data_file = f"data/questions-{subject.lower()}.json"
    if not os.path.exists(data_file):
        print(f"Error: {data_file} not found.")
        return False
        
    clean_subject_name = subject.replace('-', '_').title()
    subject_display_name = subject.replace('-', ' ').title()
    
    os.makedirs(output_dir, exist_ok=True)
    out_filename = f"JAMB_{clean_subject_name}_Questions.docx"
    out_path = os.path.join(output_dir, out_filename)
    
    print(f"\n=======================================================")
    print(f"Exporting: {subject_display_name} -> {out_path}")
    print(f"=======================================================")
    
    with open(data_file, 'r', encoding='utf-8') as f:
        data_obj = json.load(f)
        
    questions = data_obj.get('questions', [])
    print(f"Total loaded questions: {len(questions)}")
    
    # Group by questionYear
    years_dict = defaultdict(list)
    for q in questions:
        y = q.get('questionYear') or q.get('year') or 0
        try:
            y = int(y)
        except (ValueError, TypeError):
            y = 0
        years_dict[y].append(q)
        
    # Sort years descending (recent first: 2025 -> 1978)
    sorted_years = sorted(years_dict.keys(), reverse=True)
    if 0 in sorted_years:
        sorted_years.remove(0)
        sorted_years.append(0)  # Unknown at the end
        
    print(f"Years found: {len(sorted_years)} (from {sorted_years[0]} down to {sorted_years[-1]})")
    
    # Initialize Document
    doc = Document()
    
    # Set page margins (0.75 in all around for clean printable layout)
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)
        
    # Document Cover / Title Header
    doc_title_p = doc.add_paragraph()
    doc_title_p.paragraph_format.space_before = Pt(12)
    doc_title_p.paragraph_format.space_after = Pt(4)
    doc_title_p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    title_run = doc_title_p.add_run(f"JAMB UTME PAST QUESTIONS")
    format_run(title_run, font_name="Calibri", size_pt=22, bold=True, color_rgb=RGBColor(0x1A, 0x36, 0x5D))
    
    sub_title_p = doc.add_paragraph()
    sub_title_p.paragraph_format.space_before = Pt(0)
    sub_title_p.paragraph_format.space_after = Pt(4)
    sub_title_p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    sub_run = sub_title_p.add_run(f"SUBJECT: {subject_display_name.upper()}  |  YEARS: {sorted_years[-1]} – {sorted_years[0]}")
    format_run(sub_run, font_name="Calibri", size_pt=12, bold=True, color_rgb=RGBColor(0x2B, 0x6C, 0xB0))
    
    edition_p = doc.add_paragraph()
    edition_p.paragraph_format.space_before = Pt(0)
    edition_p.paragraph_format.space_after = Pt(20)
    edition_p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    ed_run = edition_p.add_run("Teacher & Mock Exam Edition  •  Questions & Yearly Answer Keys")
    format_run(ed_run, font_name="Calibri", size_pt=10, italic=True, color_rgb=RGBColor(0x71, 0x80, 0x96))
    
    # Divider line
    div_p = doc.add_paragraph()
    div_p.paragraph_format.space_after = Pt(16)
    pBrd = parse_xml(f'<w:pBrd {nsdecls("w")}><w:bottom w:val="single" w:sz="12" w:space="1" w:color="1A365D"/></w:pBrd>')
    div_p._p.get_or_add_pPr().append(pBrd)

    total_images_embedded = 0
    
    # Loop over each year
    for yr_idx, year in enumerate(sorted_years):
        year_label = str(year) if year != 0 else "Special / Additional Questions"
        year_qs = years_dict[year]
        
        # Year Section Header
        yr_head = doc.add_paragraph()
        yr_head.paragraph_format.space_before = Pt(14)
        yr_head.paragraph_format.space_after = Pt(6)
        yr_head.paragraph_format.keep_with_next = True
        
        # Header bottom border
        yrBrd = parse_xml(f'<w:pBrd {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="2" w:color="2B6CB0"/></w:pBrd>')
        yr_head._p.get_or_add_pPr().append(yrBrd)
        
        y_run = yr_head.add_run(f"EXAMINATION YEAR: {year_label}")
        format_run(y_run, font_name="Calibri", size_pt=15, bold=True, color_rgb=RGBColor(0x1A, 0x36, 0x5D))
        
        y_count_run = yr_head.add_run(f"   ({len(year_qs)} Questions)")
        format_run(y_count_run, font_name="Calibri", size_pt=10.5, italic=True, color_rgb=RGBColor(0x71, 0x80, 0x96))
        
        # Render Questions for this year
        for q_idx, q in enumerate(year_qs):
            q_num = q_idx + 1
            
            # Question stem paragraph
            q_p = doc.add_paragraph()
            q_p.paragraph_format.left_indent = Inches(0.32)
            q_p.paragraph_format.first_line_indent = Inches(-0.32)
            q_p.paragraph_format.space_before = Pt(7)
            q_p.paragraph_format.space_after = Pt(3)
            q_p.paragraph_format.line_spacing = 1.15
            
            # Bold Question Number
            q_num_run = q_p.add_run(f"{q_num}.  ")
            format_run(q_num_run, font_name="Calibri", size_pt=10.5, bold=True, color_rgb=RGBColor(0x1A, 0x20, 0x2C))
            
            # Cleaned Question Text
            cleaned_stem = clean_text(q.get('text', ''))
            stem_lines = cleaned_stem.split('\n')
            for s_idx, s_line in enumerate(stem_lines):
                s_line = s_line.strip()
                if not s_line:
                    continue
                if s_idx > 0:
                    q_p.add_run().add_break()
                line_run = q_p.add_run(s_line)
                format_run(line_run, font_name="Calibri", size_pt=10.5, color_rgb=RGBColor(0x2D, 0x37, 0x48))
                
            # Embedded Diagram / Image (if present)
            img_url = q.get('image_url') or q.get('imageUrl') or q.get('image')
            if img_url:
                clean_img_path = img_url.lstrip('/')
                if os.path.exists(clean_img_path):
                    try:
                        with Image.open(clean_img_path) as im:
                            w_px, h_px = im.size
                            
                        # Scale image sensibly (max 4.0 in wide, max 3.2 in high)
                        max_w = 4.0
                        max_h = 3.2
                        w_in = w_px / 96.0
                        h_in = h_px / 96.0
                        scale = min(max_w / w_in, max_h / h_in, 1.0)
                        final_w = w_in * scale
                        
                        img_p = doc.add_paragraph()
                        img_p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
                        img_p.paragraph_format.space_before = Pt(4)
                        img_p.paragraph_format.space_after = Pt(5)
                        img_p.paragraph_format.keep_with_next = True
                        
                        img_run = img_p.add_run()
                        try:
                            img_run.add_picture(clean_img_path, width=Inches(final_w))
                            total_images_embedded += 1
                        except Exception:
                            # Re-encode in memory with PIL to fix non-standard/truncated headers
                            import io
                            with Image.open(clean_img_path) as im_fb:
                                buf = io.BytesIO()
                                im_fb.save(buf, format='PNG')
                                buf.seek(0)
                                img_run.add_picture(buf, width=Inches(final_w))
                                total_images_embedded += 1
                    except Exception as img_err:
                        print(f"Warning embedding {clean_img_path}: {img_err}")

            # Multiple Choice Options
            opts = q.get('options', [])
            for opt_idx, opt_text in enumerate(opts):
                opt_letter = OPTIONS_LETTERS[opt_idx] if opt_idx < len(OPTIONS_LETTERS) else f"({opt_idx+1})"
                cleaned_opt = clean_text(str(opt_text))
                
                opt_p = doc.add_paragraph()
                opt_p.paragraph_format.left_indent = Inches(0.58)
                opt_p.paragraph_format.first_line_indent = Inches(-0.26)
                opt_p.paragraph_format.space_before = Pt(0)
                opt_p.paragraph_format.space_after = Pt(2)
                opt_p.paragraph_format.line_spacing = 1.15
                
                # Bold Option letter
                opt_lbl_run = opt_p.add_run(f"{opt_letter}. ")
                format_run(opt_lbl_run, font_name="Calibri", size_pt=10, bold=True, color_rgb=RGBColor(0x2B, 0x6C, 0xB0))
                
                # Option text
                opt_val_run = opt_p.add_run(cleaned_opt)
                format_run(opt_val_run, font_name="Calibri", size_pt=10, color_rgb=RGBColor(0x2D, 0x37, 0x48))
                
        # Year Answer Key Table
        add_answer_key_table(doc, year_qs, subject, year_label)
        
        # Page break between years (except after the very last year)
        if yr_idx < len(sorted_years) - 1:
            doc.add_page_break()

    print(f"Saving Word document: {out_path}...")
    doc.save(out_path)
    file_size_mb = os.path.getsize(out_path) / (1024 * 1024)
    print(f"DONE! Generated {out_filename}: {file_size_mb:.2f} MB, {total_images_embedded} images embedded.")
    return True


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Export JAMB past questions to Word docx")
    parser.add_argument("--subjects", nargs="+", default=["biology"], help="Subjects to export")
    parser.add_argument("--all", action="store_true", help="Export all 5 core subjects")
    parser.add_argument("--output_dir", type=str, default="exports", help="Output directory")
    args = parser.parse_args()
    
    if args.all:
        targets = ["biology", "mathematics", "chemistry", "physics", "english"]
    else:
        targets = args.subjects
        
    for sub in targets:
        export_subject_docx(sub, args.output_dir)
